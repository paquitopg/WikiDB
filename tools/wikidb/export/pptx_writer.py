"""PowerPoint export of a measure sheet (U13; R27).

PowerPoint does not reflow text across slides, so the writer budgets each slide by characters and continues a section
on "(suite)" slides: nothing is truncated.
"""

from __future__ import annotations

import io
import textwrap
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from .html import Paragraph, Run, Table
from .sheet import NOT_FILLED, PROVENANCE_COLUMNS, Sheet

SLIDE_BUDGET = 900          # characters of body text per slide at 16 pt
PROVENANCE_ROWS = 8         # provenance rows per slide
TITLE_LAYOUT, CONTENT_LAYOUT, TITLE_ONLY_LAYOUT = 0, 1, 5    # default template layouts


def _lines(blocks) -> list[tuple[list[Run], int, bool]]:
    """Flatten blocks into (runs, level, bullet) lines; tables become one line per row."""
    lines = []
    for block in blocks:
        if isinstance(block, Table):
            lines += [([Run(" | ".join(row))], 0, False) for row in block.rows]
        elif isinstance(block, Paragraph):
            lines.append((block.runs, block.level, bool(block.list_kind)))
    return lines


def _split(line, budget: int):
    """Split one over-long line into pieces under the budget. Whole runs are kept together where they fit; only a run
    longer than the budget is wrapped, and every piece keeps its own run's bold, italics and link."""
    runs, level, bullet = line
    if len("".join(r.text for r in runs)) <= budget:
        return [line]
    pieces, piece, used = [], [], 0
    for run in runs:
        chunks = [run] if len(run.text) <= budget else [
            Run(chunk, run.bold, run.italic, run.href)
            for chunk in textwrap.wrap(run.text, budget, break_long_words=True, break_on_hyphens=False)]
        for chunk in chunks:
            if piece and used + len(chunk.text) > budget:
                pieces.append(piece)
                piece, used = [], 0
            piece.append(chunk)
            used += len(chunk.text)
    if piece:
        pieces.append(piece)
    return [(piece, level, bullet) for piece in pieces]


def _pages(lines, budget: int = SLIDE_BUDGET):
    pages, page, used = [], [], 0
    for line in lines:
        for piece in _split(line, budget):
            size = len("".join(r.text for r in piece[0])) + 40     # a line costs at least a short line's height
            if page and used + size > budget:
                pages.append(page)
                page, used = [], 0
            page.append(piece)
            used += size
    if page:
        pages.append(page)
    return pages or [[([Run(NOT_FILLED, italic=True)], 0, False)]]


def _content_slide(prs, title: str):
    slide = prs.slides.add_slide(prs.slide_layouts[CONTENT_LAYOUT])
    slide.shapes.title.text = title
    return slide


def _fill(slide, page):
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, (runs, level, bullet) in enumerate(page):
        para = body.paragraphs[0] if i == 0 else body.add_paragraph()
        para.level = min(level + (1 if bullet else 0), 4)
        for r in runs:
            run = para.add_run()
            run.text = r.text
            run.font.size = Pt(16)
            run.font.bold, run.font.italic = r.bold or None, r.italic or None
            if r.href:
                run.hyperlink.address = r.href


def _section(prs, title: str, blocks):
    for n, page in enumerate(_pages(_lines(blocks))):
        _fill(_content_slide(prs, title if n == 0 else f"{title} (suite)"), page)


def build(sheet: Sheet, template: Path | None = None) -> bytes:
    use_template = bool(template and template.exists())
    prs = Presentation(str(template)) if use_template else Presentation()
    if not use_template:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = f"{sheet.identifier} - {sheet.title}"

    first = prs.slides.add_slide(prs.slide_layouts[TITLE_LAYOUT])
    first.shapes.title.text = sheet.title
    facts = [f"{sheet.identifier} · {sheet.bureau} · {sheet.perimeter}"]
    for name in ("dateCreation", "vecteurJuridique"):
        if name in sheet.fields and sheet.fields[name].text:
            facts.append(f"{sheet.fields[name].label} : {sheet.fields[name].text}")
    facts.append(f"Dernière modification : {sheet.last_modified}")
    first.placeholders[1].text_frame.text = "\n".join(facts)

    for name in ("raisonEtre", "modalitesModification"):
        if name in sheet.fields:
            _section(prs, sheet.fields[name].label, sheet.fields[name].blocks)
    documentation = []
    for name in ("documentationInterne", "documentationExterne"):
        if name in sheet.fields:
            documentation.append(Paragraph([Run(sheet.fields[name].label, bold=True)]))
            documentation += sheet.fields[name].blocks or [Paragraph([Run(NOT_FILLED, italic=True)])]
    _section(prs, "Documentation", documentation)

    rows = sheet.provenances or [None]
    for start in range(0, len(rows), PROVENANCE_ROWS):
        title = "Provenance des données" + (" (suite)" if start else "")
        slide = prs.slides.add_slide(prs.slide_layouts[TITLE_ONLY_LAYOUT])
        slide.shapes.title.text = title
        chunk = rows[start:start + PROVENANCE_ROWS]
        if chunk == [None]:
            box = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), prs.slide_width - Inches(1.2), Inches(1))
            box.text_frame.text = "Aucune provenance renseignée."
            continue
        shape = slide.shapes.add_table(len(chunk) + 1, len(PROVENANCE_COLUMNS), Inches(0.4), Inches(1.5),
                                       prs.slide_width - Inches(0.8), Inches(0.4) * (len(chunk) + 1))
        table = shape.table
        for j, (_, label) in enumerate(PROVENANCE_COLUMNS):
            table.cell(0, j).text = label
        for i, p in enumerate(chunk, start=1):
            for j, (key, _) in enumerate(PROVENANCE_COLUMNS):
                cell = table.cell(i, j)
                cell.text = p.values[key]
                if key == "emetteur" and p.sender_url and p.values[key]:
                    cell.text_frame.paragraphs[0].runs[0].hyperlink.address = p.sender_url
                for para in cell.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(12)

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
