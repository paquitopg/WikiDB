"""Word export of a measure sheet (U13; R27)."""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .html import Paragraph, Run, Table
from .sheet import IDENTIFICATION, PROVENANCE_COLUMNS, SECTIONS, Sheet


def _hyperlink(paragraph, url: str, text: str, bold=False, italic=False):
    rel = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    props.append(style)
    for flag, tag in ((bold, "w:b"), (italic, "w:i")):
        if flag:
            props.append(OxmlElement(tag))
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _runs(paragraph, runs: list[Run]):
    for r in runs:
        if r.href:
            _hyperlink(paragraph, r.href, r.text, r.bold, r.italic)
        else:
            run = paragraph.add_run(r.text)
            run.bold, run.italic = r.bold or None, r.italic or None


def _style(document, name: str, fallback: str = "Normal") -> str:
    return name if name in [s.name for s in document.styles] else fallback


def _blocks(document, blocks):
    if not blocks:
        document.add_paragraph().add_run("Non renseigné.").italic = True
        return
    for block in blocks:
        if isinstance(block, Table):
            _table(document, block.rows, block.header)
            continue
        style = "Normal"
        if block.heading:
            style = _style(document, "Heading 3")
        elif block.list_kind:
            base = "List Bullet" if block.list_kind == "bullet" else "List Number"
            style = _style(document, base + (f" {block.level + 1}" if block.level else ""), _style(document, base))
        _runs(document.add_paragraph(style=style), block.runs)


def _table(document, rows: list[list[str]], header: bool):
    width = max(len(r) for r in rows)
    table = document.add_table(rows=0, cols=width)
    table.style = _style(document, "Table Grid", table.style.name if table.style else "Normal")
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cells[j].text = value
            if header and i == 0:
                for run in cells[j].paragraphs[0].runs:
                    run.bold = True
    return table


def build(sheet: Sheet, template: Path | None = None) -> bytes:
    document = Document(str(template)) if template and template.exists() else Document()
    document.core_properties.title = f"{sheet.identifier} - {sheet.title}"
    document.add_heading(sheet.title, level=0)

    ident = [["Identifiant", sheet.identifier], ["Bureau propriétaire", sheet.bureau], ["Périmètre", sheet.perimeter]]
    ident += [[sheet.fields[n].label, sheet.fields[n].text or "Non renseigné"] for n in IDENTIFICATION
              if n in sheet.fields]
    _table(document, ident, header=False)

    for name in SECTIONS:
        if name in sheet.fields:
            document.add_heading(sheet.fields[name].label, level=1)
            _blocks(document, sheet.fields[name].blocks)

    document.add_heading("Provenance des données", level=1)
    if sheet.provenances:
        table = _table(document, [[label for _, label in PROVENANCE_COLUMNS]], header=True)
        for p in sheet.provenances:
            cells = table.add_row().cells
            for j, (key, _) in enumerate(PROVENANCE_COLUMNS):
                if key == "emetteur" and p.sender_url and p.values[key]:
                    _hyperlink(cells[j].paragraphs[0], p.sender_url, p.values[key])
                else:
                    cells[j].text = p.values[key]
    else:
        document.add_paragraph().add_run("Aucune provenance renseignée.").italic = True

    footer = document.add_paragraph()
    footer.add_run(f"Dernière modification : {sheet.last_modified}. Fiche en ligne : ")
    _hyperlink(footer, sheet.url, sheet.url)

    out = io.BytesIO()
    document.save(out)
    return out.getvalue()
