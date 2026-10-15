import io

from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation

from wikidb.export import docx_writer, pptx_writer
from wikidb.export.html import absolute_url
from wikidb.export.server import create_app
from wikidb.export.sheet import from_export_data

ORIGIN = "https://wiki.example.org"
INTERNAL = "http://wikidb:8080"


def export_data(**overrides):
    fields = {
        "cleSource": "K1", "dateCreation": "01/02/2020", "vecteurJuridique": "LFI 2020, article 1",
        "raisonEtre": '<p>Soutenir <strong>l\'activité</strong>, voir <a href="/bin/view/Annuaire/Insee">l\'Insee</a>.</p>'
                      "<ul><li>premier point</li><li>second point</li></ul>",
        "modalitesModification": "<p>Par loi de finances.</p>",
        "documentationInterne": "", "documentationExterne": "<table><tr><th>Texte</th></tr><tr><td>Rapport</td></tr></table>",
    }
    data = {
        "identifier": "M-0042", "title": "Prime d'activité", "url": "/bin/view/Mesures/A/Secteur/M-0042",
        "bureau": {"code": "A", "name": "Bureau A"}, "perimetre": {"code": "Secteur", "label": "Sectoriel"},
        "fields": [{"name": k, "label": k.upper(), "html": v} for k, v in fields.items()],
        "provenances": [{"donnee": "Effectifs", "emetteur": "Insee", "emetteurUrl": "/bin/view/Annuaire/Insee",
                         "frequence": "Mensuelle", "format": "Fichier Excel", "contact": "dp@example.org"},
                        {"donnee": "Crédits", "emetteur": "DGFiP", "emetteurUrl": "", "frequence": "",
                         "format": "", "contact": ""}],
        "lastModified": {"date": "15/10/2026 09:00", "author": "Agent A"},
    }
    data.update(overrides)
    return data


def sheet(**overrides):
    return from_export_data(export_data(**overrides), ORIGIN, (INTERNAL,))


def docx_text(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    parts += [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return "\n".join(parts)


def docx_links(content: bytes) -> set[str]:
    doc = Document(io.BytesIO(content))
    return {r.target_ref for r in doc.part.rels.values() if r.reltype.endswith("/hyperlink")}


def slides_text(content: bytes) -> list[str]:
    prs = Presentation(io.BytesIO(content))
    out = []
    for slide in prs.slides:
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if shape.has_table:
                texts += [c.text for row in shape.table.rows for c in row.cells]
        out.append("\n".join(texts))
    return out


def test_internal_links_move_to_the_public_origin():
    assert absolute_url("/bin/view/A", ORIGIN, (INTERNAL,)) == f"{ORIGIN}/bin/view/A"
    assert absolute_url(f"{INTERNAL}/bin/view/A?x=1", ORIGIN, (INTERNAL,)) == f"{ORIGIN}/bin/view/A?x=1"
    assert absolute_url("http://localhost:8080/bin/view/A", ORIGIN) == f"{ORIGIN}/bin/view/A"
    assert absolute_url("https://www.legifrance.gouv.fr/x", ORIGIN, (INTERNAL,)) == "https://www.legifrance.gouv.fr/x"


def test_word_document_holds_every_field_the_provenance_table_and_working_links():
    content = docx_writer.build(sheet())
    text = docx_text(content)
    for expected in ("Prime d'activité", "M-0042", "A – Bureau A", "Sectoriel", "LFI 2020, article 1", "01/02/2020",
                     "Soutenir", "premier point", "Par loi de finances.", "Rapport", "Effectifs", "Fichier Excel",
                     "DGFiP", "15/10/2026 09:00 par Agent A"):
        assert expected in text, expected
    assert f"{ORIGIN}/bin/view/Annuaire/Insee" in docx_links(content)
    styles = {p.style.name for p in Document(io.BytesIO(content)).paragraphs if "point" in p.text}
    assert styles == {"List Bullet"}


def test_word_marks_empty_fields_instead_of_dropping_them():
    assert "Non renseigné." in docx_text(docx_writer.build(sheet()))


def test_presentation_has_a_title_slide_and_one_provenance_row_per_block():
    slides = slides_text(pptx_writer.build(sheet()))
    assert slides[0].startswith("Prime d'activité")
    provenance = [s for s in slides if s.startswith("Provenance des données")]
    assert len(provenance) == 1 and "Effectifs" in provenance[0] and "DGFiP" in provenance[0]


def test_long_rationale_continues_on_follow_on_slides_without_losing_text():
    words = [f"mot{i:04d}" for i in range(800)]                     # about 6,400 characters
    data = export_data()
    data["fields"][3]["html"] = "<p>" + " ".join(words) + "</p>"
    slides = slides_text(pptx_writer.build(from_export_data(data, ORIGIN)))
    rationale = [s for s in slides if s.startswith("RAISONETRE")]
    assert len(rationale) >= 6
    assert rationale[1].startswith("RAISONETRE (suite)")
    joined = " ".join(rationale)
    assert all(w in joined for w in words)


def test_file_names_are_valid_on_windows():
    assert sheet(title='Taux "réduit" : 5/10 ?').file_name("docx") == "M-0042 - Taux -réduit- - 5-10 -.docx"


def client(status, data=None, seen=None):
    def fetch(internal, reference, headers):
        if seen is not None:
            seen.update(reference=reference, headers=headers)
        return status, data
    return TestClient(create_app(internal_url=INTERNAL, fetch=fetch), follow_redirects=False)


def test_export_forwards_the_reader_credentials_and_returns_the_file():
    seen = {}
    r = client(200, export_data(), seen).get("/export/Mesures/A/Secteur/M-0042.docx",
                                              headers={"Cookie": "JSESSIONID=abc", "Authorization": "Basic eA=="})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml")
    assert "filename*=UTF-8''M-0042%20-%20Prime%20d%27activit%C3%A9.docx" in r.headers["content-disposition"]
    assert seen == {"reference": "Mesures.A.Secteur.M-0042",
                    "headers": {"cookie": "JSESSIONID=abc", "authorization": "Basic eA=="}}


def test_redirect_to_login_becomes_401_with_a_sign_in_link():
    r = client(302).get("/export/Mesures/A/Secteur/M-0042.pptx")
    assert r.status_code == 401
    assert "/bin/login/XWiki/XWikiLogin?xredirect=/export/Mesures/A/Secteur/M-0042.pptx" in r.text


def test_denied_or_missing_sheet_gets_a_404_that_names_nothing():
    for status in (403, 404):
        r = client(status).get("/export/Mesures/B/Secteur/M-0007.docx")
        assert r.status_code == 404 and "M-0007" not in r.text


def test_paths_outside_measure_sheets_are_refused_without_calling_the_wiki():
    seen = {}
    for path in ("/export/XWiki/XWikiPreferences.docx", "/export/Mesures/A/Secteur/M-0042.pdf",
                 "/export/Bureaux/A/WebHome.docx"):
        assert client(200, export_data(), seen).get(path).status_code == 404
    assert seen == {}
