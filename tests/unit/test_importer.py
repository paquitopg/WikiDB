from pathlib import Path

import pytest
from openpyxl import Workbook

from fakes import FakeXWiki
from wikidb.importer import ImportRejected, execute, plan, read_rows, validate
from wikidb.measures import MESURE_CLASS
from wikidb.provision import Bureau

BUREAUX = [Bureau("A", "Bureau A", "secteur"), Bureau("S1", "Synthèse 1", "synthese")]
HEADER = "bureau;perimetre;intitule;cle_source;raison_etre\n"


def write(tmp_path: Path, text: str) -> Path:
    path = tmp_path / "mesures.csv"
    path.write_text(text, encoding="utf-8")
    return path


def run(wiki, path, dry_run=False):
    the_plan = plan(wiki, validate(read_rows(path), BUREAUX))
    if not dry_run:
        execute(wiki, the_plan)
    return the_plan


def test_valid_rows_create_sheets_in_their_perimeter_space_with_neutral_names(tmp_path):
    wiki = FakeXWiki()
    path = write(tmp_path, HEADER + "A;secteur;Prime d'activité;K1;Soutien\nA;synthese;Crédit d'impôt;K2;\n"
                 "S1;secteur;Mesure S;K3;\n")
    run(wiki, path)
    assert wiki.pages["Mesures.A.Secteur.M-0001"]["title"] == "Prime d'activité"
    assert wiki.pages["Mesures.A.Synthese.M-0002"]["title"] == "Crédit d'impôt"
    assert wiki.pages["Mesures.S1.Secteur.M-0003"]["title"] == "Mesure S"
    assert wiki.objs("Mesures.A.Secteur.M-0001", MESURE_CLASS)[0]["raisonEtre"] == "Soutien"


def test_rerunning_the_same_file_creates_nothing(tmp_path):
    wiki = FakeXWiki()
    path = write(tmp_path, HEADER + "A;secteur;Prime;K1;Soutien\n")
    run(wiki, path)
    second = run(wiki, path)
    assert second.create == [] and second.fill == []
    assert len([p for p in wiki.pages if p.startswith("Mesures.")]) == 1


def test_rerun_fills_empty_fields_but_keeps_what_a_champion_wrote(tmp_path):
    wiki = FakeXWiki()
    run(wiki, write(tmp_path, HEADER + "A;secteur;Prime;K1;\n"))
    obj = wiki.objects["Mesures.A.Secteur.M-0001"][MESURE_CLASS][0]
    obj["raisonEtre"] = "Texte du champion"
    run(wiki, write(tmp_path, "bureau;perimetre;intitule;cle_source;raison_etre;vecteur_juridique\n"
                              "A;secteur;Prime;K1;Texte importé;LFI 2019\n"))
    assert obj["raisonEtre"] == "Texte du champion"
    assert obj["vecteurJuridique"] == "LFI 2019"


@pytest.mark.parametrize("line, message", [
    ("Z;secteur;Prime;K1;", "unknown bureau code 'Z'"),
    ("A;autre;Prime;K1;", "perimeter must be"),
    ("A;secteur;;K1;", "empty title"),
])
def test_invalid_row_rejects_the_file_and_writes_nothing(tmp_path, line, message):
    wiki = FakeXWiki()
    with pytest.raises(ImportRejected, match=message):
        run(wiki, write(tmp_path, HEADER + "A;secteur;Valide;K0;\n" + line + "\n"))
    assert wiki.pages == {}


def test_dry_run_writes_nothing_and_lists_planned_actions(tmp_path):
    wiki = FakeXWiki()
    the_plan = run(wiki, write(tmp_path, HEADER + "A;secteur;Prime;K1;\n"), dry_run=True)
    assert wiki.pages == {}
    assert the_plan.describe() == ["create Mesures.A.Secteur.M-0001 « Prime » (line 2)"]


def test_xlsx_with_accents_and_merged_header_cells(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.append(["Bureau", "Périmètre", "Intitulé", "Clé source", None])
    sheet.merge_cells("D1:E1")
    sheet.append(["A", "secteur", "Réduction d'impôt « dons »", "K9", None])
    path = tmp_path / "mesures.xlsx"
    book.save(path)
    wiki = FakeXWiki()
    run(wiki, path)
    assert wiki.pages["Mesures.A.Secteur.M-0001"]["title"] == "Réduction d'impôt « dons »"
