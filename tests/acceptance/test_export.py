"""Word and PowerPoint export against the dev wiki (U13; R7, R22, R27, AE10).

The export service runs in-process on a free local port and reaches the dev wiki at --base-url, exactly as the
deployed service reaches XWiki: it forwards the agent's credentials and holds none of its own.
"""

import io
import socket
import threading
import time
from urllib.parse import urlsplit

import pytest
import requests
import uvicorn
from docx import Document
from pptx import Presentation

from test_measures import FILLED, client_for, create, created_page
from wikidb.client import PageRef
from wikidb.export.server import create_app
from wikidb.measures import FIELDS, MESURE_CLASS, PROVENANCE_CLASS


@pytest.fixture(scope="module")
def service(base_url):
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        port = s.getsockname()[1]
    server = uvicorn.Server(uvicorn.Config(create_app(internal_url=base_url), host="127.0.0.1", port=port,
                                           log_level="warning"))
    thread = threading.Thread(target=server.run, daemon=True)
    thread.start()
    deadline = time.time() + 20
    while not server.started and time.time() < deadline:
        time.sleep(0.1)
    yield f"http://127.0.0.1:{port}"
    server.should_exit = True
    thread.join(timeout=10)


@pytest.fixture(scope="module")
def sheets(world, robot):
    """A complete sheet of A with two provenance blocks, and a sector sheet of B."""
    a = world["agents"]["test-agent-a"]
    own = created_page(create(a, "Mesure exportée zorglubX", "A"))
    sender = PageRef(("Annuaire", "Externes"), "TestEmetteurExport")
    robot.put_page(sender, "Émetteur d'export")
    writer = client_for(a)
    writer.update_object(own, MESURE_CLASS, 0, FILLED)
    writer.add_object(own, PROVENANCE_CLASS, {"donnee": "Effectifs exportés", "emetteur": "Service émetteur",
                                              "ficheAnnuaire": "Annuaire.Externes.TestEmetteurExport",
                                              "frequence": "mensuelle", "format": "Fichier Excel"})
    writer.add_object(own, PROVENANCE_CLASS, {"donnee": "Crédits exportés", "emetteur": "Autre émetteur"})
    other = PageRef(("Mesures", "B", "Secteur"), "M-9201")
    robot.put_page(other, "Mesure secrète de B zorglubY")
    if not robot.list_objects(other, MESURE_CLASS):
        robot.add_object(other, MESURE_CLASS, {f: "" for f in FIELDS})
    yield {"own": own, "other": other}
    for page in (own, other, sender):
        robot.delete_page(page)


def export_url(service: str, page: PageRef, extension: str) -> str:
    return f"{service}/export/{'/'.join(page.spaces)}/{page.name}.{extension}"


def get(agent, url: str, **kwargs):
    return requests.get(url, auth=agent.session.auth if agent.name else None, allow_redirects=False, timeout=120,
                        **kwargs)


def docx_text(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    return "\n".join([p.text for p in doc.paragraphs] +
                     [c.text for t in doc.tables for row in t.rows for c in row.cells])


def test_ae10_agent_of_another_sector_bureau_gets_no_file_and_no_title(world, service, sheets):
    a = world["agents"]["test-agent-a"]
    for extension in ("docx", "pptx"):
        r = get(a, export_url(service, sheets["other"], extension))
        assert r.status_code == 404
        assert "zorglubY" not in r.text and "M-9201" not in r.text


def test_synthesis_agent_exports_a_sector_sheet_of_another_bureau(world, service, sheets):
    r = get(world["agents"]["test-agent-s1"], export_url(service, sheets["other"], "docx"))
    assert r.status_code == 200
    assert "zorglubY" in docx_text(r.content)


def test_guest_and_expired_session_get_401_without_following_the_login_redirect(world, service, sheets):
    url = export_url(service, sheets["own"], "docx")
    assert get(world["agents"]["guest"], url).status_code == 401
    expired = requests.get(url, headers={"Cookie": "JSESSIONID=expired"}, allow_redirects=False, timeout=60)
    assert expired.status_code == 401 and "Se connecter" in expired.text


def test_word_export_of_a_complete_sheet_holds_its_content_and_public_links(world, service, sheets, base_url):
    a = world["agents"]["test-agent-a"]
    public = urlsplit(base_url)
    r = get(a, export_url(service, sheets["own"], "docx"),
            headers={"X-Forwarded-Proto": public.scheme, "X-Forwarded-Host": public.netloc})
    assert r.status_code == 200
    text = docx_text(r.content)
    for expected in [*FILLED.values(), "Effectifs exportés", "Crédits exportés", "Bureau fictif A", "Sectoriel",
                     "zorglubX", "Mensuelle"]:
        assert expected in text, expected
    links = {rel.target_ref for rel in Document(io.BytesIO(r.content)).part.rels.values()
             if rel.reltype.endswith("/hyperlink")}
    assert f"{public.scheme}://{public.netloc}/bin/view/Annuaire/Externes/TestEmetteurExport" in links


def test_powerpoint_export_opens_with_title_and_provenance_rows(world, service, sheets):
    r = get(world["agents"]["test-agent-a"], export_url(service, sheets["own"], "pptx"))
    assert r.status_code == 200
    prs = Presentation(io.BytesIO(r.content))
    slides = list(prs.slides)
    assert "zorglubX" in slides[0].shapes.title.text
    tables = [s.table for slide in slides for s in slide.shapes if s.has_table]
    assert len(tables) == 1 and len(tables[0].rows) == 3                  # header + two blocks


def test_export_data_answers_404_for_anything_but_a_readable_measure_sheet(world):
    a = world["agents"]["test-agent-a"]
    for reference in ("XWiki.XWikiPreferences", "Bureaux.B.WebHome", "Mesures.A.Secteur.M-0000",
                      "Mesures.B.Secteur.M-9201"):
        r = a.session.get(f"{a.base_url}/bin/get/Mesures/Code/ExportData",
                          params={"xpage": "plain", "outputSyntax": "plain", "mesure": reference},
                          headers={"Accept": "application/json"})
        assert r.status_code == 404, reference
        assert "zorglubY" not in r.text


def test_code_in_a_sheet_field_does_not_run_in_the_export(world, service, sheets):
    a = world["agents"]["test-agent-a"]
    client_for(a).update_object(sheets["own"], MESURE_CLASS, 0,
                                {"raisonEtre": "{{velocity}}RESULTAT-$mathtool.add(40, 2){{/velocity}}"})
    try:
        r = get(a, export_url(service, sheets["own"], "docx"))
        assert r.status_code == 200
        assert "RESULTAT-42" not in docx_text(r.content)
    finally:
        client_for(a).update_object(sheets["own"], MESURE_CLASS, 0, {"raisonEtre": FILLED["raisonEtre"]})
